"""
Excel Import / Export Engine for Gestione Ore Dipendenti.
Generates full 3-sheet professional workbooks and parses existing sheets.
"""

import os
import csv
import datetime
from typing import Optional, Dict, Any, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import database


def export_to_excel(employee_id: int, output_path: str) -> str:
    """Exports employee data to a rich 3-sheet Excel workbook."""
    emp = database.get_employee(employee_id)
    if not emp:
        raise ValueError(f"Dipendente con ID {employee_id} non trovato.")

    dash_data = database.get_dashboard_data(employee_id)
    summary = dash_data["summary"]
    monthly_stats = dash_data["monthly_stats"]
    entries = database.get_all_chronological_entries(employee_id)

    wb = openpyxl.Workbook()

    # 1. Sheet: Riepilogo & Dashboard
    ws_dash = wb.active
    ws_dash.title = "Riepilogo & Dashboard"
    ws_dash.views.sheetView[0].showGridLines = True

    # 2. Sheet: Registro Movimenti
    ws_log = wb.create_sheet(title="Registro Movimenti")
    ws_log.views.sheetView[0].showGridLines = True

    # 3. Sheet: Vista Calendario Mensile
    ws_cal = wb.create_sheet(title="Vista Calendario Mensile")
    ws_cal.views.sheetView[0].showGridLines = True

    # Palette
    NAVY = "1B365D"
    DARK_SLATE = "2C3E50"
    DARK_TEAL = "134E6F"
    ACCENT_BLUE = "2E75B6"
    LIGHT_GRAY = "F9FAFB"
    BORDER_GRAY = "D3D3D3"
    WHITE = "FFFFFF"
    BLACK = "000000"
    COLOR_YELLOW = "FFFF00"
    COLOR_GREEN = "00FF00"

    fill_header = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
    fill_accent = PatternFill(start_color=ACCENT_BLUE, end_color=ACCENT_BLUE, fill_type="solid")
    fill_sub = PatternFill(start_color=DARK_SLATE, end_color=DARK_SLATE, fill_type="solid")
    fill_zebra = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")
    fill_yellow = PatternFill(start_color=COLOR_YELLOW, end_color=COLOR_YELLOW, fill_type="solid")
    fill_green = PatternFill(start_color=COLOR_GREEN, end_color=COLOR_GREEN, fill_type="solid")
    fill_tot = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    fill_info = PatternFill(start_color="EBF1F5", end_color="EBF1F5", fill_type="solid")
    fill_kpi = PatternFill(start_color="F2F7FA", end_color="F2F7FA", fill_type="solid")

    font_title = Font(name="Calibri", size=15, bold=True, color=WHITE)
    font_sec_header = Font(name="Calibri", size=11, bold=True, color=WHITE)
    font_tbl_header = Font(name="Calibri", size=10, bold=True, color=WHITE)
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_regular = Font(name="Calibri", size=10)
    font_big_num = Font(name="Calibri", size=24, bold=True, color=BLACK)
    font_box_label = Font(name="Calibri", size=10, bold=True, color="262626")
    font_box_num = Font(name="Calibri", size=15, bold=True, color=BLACK)

    thin_side = Side(style='thin', color=BORDER_GRAY)
    med_side = Side(style='medium', color=NAVY)
    double_side = Side(style='double', color=NAVY)
    border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    border_header = Border(left=thin_side, right=thin_side, top=med_side, bottom=med_side)
    border_total = Border(left=thin_side, right=thin_side, top=thin_side, bottom=double_side)

    # -------------------------------------------------------------
    # POPULATE SHEET 1: RIEPILOGO & DASHBOARD
    # -------------------------------------------------------------
    ws_dash.merge_cells("B1:N1")
    ws_dash["B1"] = f"QUADRO DI CONTROLLO STRAORDINARI E BANCA ORE - {emp['name'].upper()}"
    ws_dash["B1"].font = font_title
    ws_dash["B1"].fill = fill_header
    ws_dash["B1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_dash.row_dimensions[1].height = 36

    # Parametri Base (B3:C10)
    ws_dash.merge_cells("B3:C3")
    ws_dash["B3"] = "PARAMETRI BASE"
    ws_dash["B3"].font = font_sec_header
    ws_dash["B3"].fill = fill_accent
    ws_dash["B3"].alignment = Alignment(horizontal="center", vertical="center")

    params = [
        ("Ore Base Iniziali (Banca Ore)", summary["initial_bank_hours"], '#,##0.0 "h"'),
        ("Totale Ore Utilizzate (Gialle)", summary["totale_ore_utilizzate"], '#,##0.0 "h"'),
        ("Ore Straordinari Svolte (Verdi)", summary["totale_ore_maturate"], '#,##0.0 "h"'),
        ("Orario Giornaliero Standard (Ferie da 7h)", emp["standard_day_hours"], '0.0 "ore"'),
        ("Orario Giornaliero Breve (Ferie da 4h)", emp["short_day_hours"], '0.0 "ore"'),
        ("Orario Settimanale Part-time (da 18h)", emp["part_time_week_hours"], '0.0 "ore"')
    ]

    for idx, (lbl, val, fmt) in enumerate(params, start=5):
        c_lbl = ws_dash.cell(row=idx, column=2, value=lbl)
        c_lbl.font = font_regular
        c_lbl.border = border_all
        c_val = ws_dash.cell(row=idx, column=3, value=val)
        c_val.font = font_bold
        c_val.number_format = fmt
        c_val.alignment = Alignment(horizontal="right")
        c_val.border = border_all
        if idx in (5, 6):
            c_val.fill = fill_zebra

    # Sintesi Movimenti & Saldo Ore (E3:I13)
    ws_dash.merge_cells("E3:I3")
    ws_dash["E3"] = "SINTESI MOVIMENTI & SALDO ORE"
    ws_dash["E3"].font = font_sec_header
    ws_dash["E3"].fill = fill_sub
    ws_dash["E3"].alignment = Alignment(horizontal="center", vertical="center")

    # Box Totale Ore Accreditate
    ws_dash.merge_cells("E5:F5")
    ws_dash["E5"] = "TOTALE ORE ACCREDITATE"
    ws_dash["E5"].font = font_box_label
    ws_dash["E5"].alignment = Alignment(horizontal="center", vertical="center")

    ws_dash.merge_cells("E6:F6")
    ws_dash["E6"] = summary["totale_ore_accreditate"]
    ws_dash["E6"].font = font_box_num
    ws_dash["E6"].alignment = Alignment(horizontal="center", vertical="center")
    ws_dash["E6"].number_format = '#,##0.00'

    # Box Totale Ore Utilizzate
    ws_dash.merge_cells("E8:F8")
    ws_dash["E8"] = "TOTALE ORE UTILIZZATE"
    ws_dash["E8"].font = font_box_label
    ws_dash["E8"].alignment = Alignment(horizontal="center", vertical="center")

    ws_dash.merge_cells("E9:F9")
    ws_dash["E9"] = -summary["totale_ore_utilizzate"]
    ws_dash["E9"].font = font_box_num
    ws_dash["E9"].alignment = Alignment(horizontal="center", vertical="center")
    ws_dash["E9"].number_format = '+#,##0.00;-#,##0.00;0.00'

    # Box ORE RIMANENTI
    ws_dash.merge_cells("H5:I5")
    ws_dash["H5"] = "ORE RIMANENTI"
    ws_dash["H5"].font = Font(name="Calibri", size=12, bold=True, color="1F4E79")
    ws_dash["H5"].fill = fill_kpi
    ws_dash["H5"].alignment = Alignment(horizontal="center", vertical="center")

    ws_dash.merge_cells("H6:I7")
    ws_dash["H6"] = summary["ore_rimanenti"]
    ws_dash["H6"].font = font_big_num
    ws_dash["H6"].fill = fill_kpi
    ws_dash["H6"].alignment = Alignment(horizontal="center", vertical="center")
    ws_dash["H6"].number_format = '#,##0.00'

    # Indicatori di Conversione
    conv_items = [
        ("Giorni extra di ferie pagati possibili (da 7 ore)", summary["ferie_7h"]),
        ("Giorni extra di ferie pagati possibili (da 4 ore)", summary["ferie_4h"]),
        ("settimane da 18 ore", summary["settimane_18h"])
    ]
    for idx, (lbl, val) in enumerate(conv_items, start=11):
        ws_dash.merge_cells(f"E{idx}:F{idx}")
        ws_dash[f"E{idx}"] = lbl
        ws_dash[f"E{idx}"].font = font_regular
        ws_dash[f"E{idx}"].border = border_all
        c_v = ws_dash.cell(row=idx, column=7, value=val)
        c_v.font = Font(name="Calibri", size=11, bold=True, color=DARK_TEAL)
        c_v.alignment = Alignment(horizontal="right")
        c_v.number_format = '#,##0.00'
        c_v.border = border_all

    # Prospetto Mensile di Dettaglio (B16:G34)
    ws_dash.merge_cells("B16:G16")
    ws_dash["B16"] = "PROSPETTO MENSILE DI DETTAGLIO (2026 - 2027)"
    ws_dash["B16"].font = font_sec_header
    ws_dash["B16"].fill = fill_header
    ws_dash["B16"].alignment = Alignment(horizontal="center", vertical="center")

    tbl_headers = ["Mese / Anno", "Ore Utilizzate (Gialle)", "Ore Straordinario (Verdi)", "Saldo Mese", "Ore Rimaste Progressive", "Stato"]
    for c_idx, h in enumerate(tbl_headers, start=2):
        c = ws_dash.cell(row=17, column=c_idx, value=h)
        c.font = font_tbl_header
        c.fill = fill_sub
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border_header

    start_row = 18
    for i, m in enumerate(monthly_stats):
        curr_r = start_row + i
        ws_dash.cell(row=curr_r, column=2, value=m["month_year"]).alignment = Alignment(horizontal="left")
        ws_dash.cell(row=curr_r, column=3, value=m["ore_utilizzate"]).number_format = '#,##0.0'
        ws_dash.cell(row=curr_r, column=4, value=m["ore_maturate"]).number_format = '#,##0.0'
        ws_dash.cell(row=curr_r, column=5, value=m["saldo_mese"]).number_format = '+#,##0.0;-#,##0.0;0.0'
        ws_dash.cell(row=curr_r, column=6, value=m["ore_rimaste_progressive"]).number_format = '#,##0.0'
        ws_dash.cell(row=curr_r, column=7, value=m["status"]).alignment = Alignment(horizontal="center")

        for c_idx in range(2, 8):
            cell = ws_dash.cell(row=curr_r, column=c_idx)
            cell.font = font_regular
            cell.border = border_all
            if i % 2 == 1:
                cell.fill = fill_zebra

    tot_r = start_row + len(monthly_stats)
    ws_dash.cell(row=tot_r, column=2, value="TOTALE COMPLESSIVO").font = font_bold
    ws_dash.cell(row=tot_r, column=3, value=summary["totale_ore_utilizzate"]).font = font_bold
    ws_dash.cell(row=tot_r, column=4, value=summary["totale_ore_maturate"]).font = font_bold
    ws_dash.cell(row=tot_r, column=5, value=-summary["totale_ore_utilizzate"]).font = font_bold
    ws_dash.cell(row=tot_r, column=6, value=summary["ore_rimanenti"]).font = font_bold

    for c_idx in range(2, 8):
        c = ws_dash.cell(row=tot_r, column=c_idx)
        c.fill = fill_tot
        c.border = border_total

    # -------------------------------------------------------------
    # POPULATE SHEET 2: REGISTRO MOVIMENTI
    # -------------------------------------------------------------
    ws_log.merge_cells("A1:I1")
    ws_log["A1"] = f"REGISTRO ANALITICO MOVIMENTI - {emp['name'].upper()}"
    ws_log["A1"].font = font_title
    ws_log["A1"].fill = fill_header
    ws_log["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_log.row_dimensions[1].height = 35

    ws_log.merge_cells("A2:I2")
    ws_log["A2"] = "Legenda: CELLE GIALLE = Ore Utilizzate (permessi presi) | CELLE VERDI = Ore Maturate (straordinari fatti)"
    ws_log["A2"].font = Font(name="Calibri", size=10, bold=True, color="1F4E79")
    ws_log["A2"].fill = fill_info
    ws_log["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws_log.row_dimensions[2].height = 22

    log_headers = [
        "N°", "Data", "Giorno", "Mese / Anno", "Tipologia / Dettaglio",
        "Ore Utilizzate (Gialle)", "Ore Maturate Straordinario (Verdi)", "Saldo Ore Residue (h)", "Note"
    ]
    for col_idx, h in enumerate(log_headers, start=1):
        c = ws_log.cell(row=3, column=col_idx, value=h)
        c.font = font_tbl_header
        c.fill = fill_sub
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border_header
    ws_log.row_dimensions[3].height = 25

    cum_util = 0.0
    for idx, e in enumerate(entries, start=1):
        r_i = 3 + idx
        ws_log.cell(row=r_i, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws_log.cell(row=r_i, column=2, value=e["date"]).alignment = Alignment(horizontal="center")
        ws_log.cell(row=r_i, column=3, value=e["day_of_week"]).alignment = Alignment(horizontal="center")
        ws_log.cell(row=r_i, column=4, value=e["month_year"]).alignment = Alignment(horizontal="center")
        course_label = f"[{e['course_name']}] " if e.get("course_name") else ""
        timing_label = f" ({e['scheduled_start']}-{e['scheduled_end']} -> {e['actual_start']}-{e['actual_end']})" if e.get("scheduled_start") and e.get("actual_end") and e.get("scheduled_start") != e.get("actual_start") or e.get("scheduled_end") != e.get("actual_end") else ""
        desc = course_label + (e["causale"] or ("Straordinario" if e["type"] == "maturata" else "Permesso")) + timing_label
        ws_log.cell(row=r_i, column=5, value=desc).alignment = Alignment(horizontal="left")

        c_u = ws_log.cell(row=r_i, column=6)
        c_m = ws_log.cell(row=r_i, column=7)

        if e["type"] == "utilizzata":
            cum_util += e["hours"]
            c_u.value = e["hours"]
            c_u.fill = fill_yellow
            c_u.number_format = '#,##0.0'
        elif e["type"] == "maturata":
            c_m.value = e["hours"]
            c_m.fill = fill_green
            c_m.number_format = '#,##0.0'

        c_u.alignment = Alignment(horizontal="right")
        c_m.alignment = Alignment(horizontal="right")

        # Saldo progressivo = Ore base - cum_util
        res_h = summary["initial_bank_hours"] - cum_util
        c_res = ws_log.cell(row=r_i, column=8, value=res_h)
        c_res.number_format = '#,##0.0'
        c_res.alignment = Alignment(horizontal="right")

        ws_log.cell(row=r_i, column=9, value=e["note"]).alignment = Alignment(horizontal="left")

        for col_i in range(1, 10):
            cell = ws_log.cell(row=r_i, column=col_i)
            cell.font = font_regular
            cell.border = border_all
            if idx % 2 == 0 and cell.fill.start_color.rgb not in (COLOR_YELLOW, COLOR_GREEN):
                cell.fill = fill_zebra

    # Totale Registro
    r_tot_log = 4 + len(entries)
    ws_log.merge_cells(f"A{r_tot_log}:E{r_tot_log}")
    ws_log.cell(row=r_tot_log, column=1, value="TOTALE COMPLESSIVO").alignment = Alignment(horizontal="right")
    ws_log.cell(row=r_tot_log, column=1).font = font_bold
    ws_log.cell(row=r_tot_log, column=6, value=summary["totale_ore_utilizzate"]).font = font_bold
    ws_log.cell(row=r_tot_log, column=7, value=summary["totale_ore_maturate"]).font = font_bold
    ws_log.cell(row=r_tot_log, column=8, value=summary["ore_rimanenti"]).font = font_bold

    for col_i in range(1, 10):
        c = ws_log.cell(row=r_tot_log, column=col_i)
        c.fill = fill_tot
        c.border = border_total

    # -------------------------------------------------------------
    # POPULATE SHEET 3: VISTA CALENDARIO MENSILE (4 MESI PER RIGA)
    # -------------------------------------------------------------
    ws_cal.merge_cells("A1:P1")
    ws_cal["A1"] = f"CALENDARIO MENSILE STRAORDINARI ({emp['name'].upper()}) - GIALLO = UTILIZZATE, VERDE = MATURATE"
    ws_cal["A1"].font = font_title
    ws_cal["A1"].fill = fill_header
    ws_cal["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_cal.row_dimensions[1].height = 35

    month_blocks = [
        (3, [("FEBBRAIO 2026", 2), ("MARZO 2026", 6), ("APRILE 2026", 10), ("MAGGIO 2026", 14)]),
        (23, [("GIUGNO 2026", 2), ("LUGLIO 2026", 6), ("AGOSTO 2026", 10), ("SETTEMBRE 2026", 14)]),
        (43, [("OTTOBRE 2026", 2), ("NOVEMBRE 2026", 6), ("DICEMBRE 2026", 10), ("GENNAIO 2027", 14)]),
        (63, [("FEBBRAIO 2027", 2), ("MARZO 2027", 6), ("APRILE 2027", 10), ("MAGGIO 2027", 14)])
    ]

    for header_r, row_months in month_blocks:
        for mname, col_start in row_months:
            # Month Header
            col_end = col_start + 2
            ws_cal.merge_cells(start_row=header_r, start_column=col_start, end_row=header_r, end_column=col_end)
            c_m = ws_cal.cell(header_r, col_start, value=mname)
            c_m.font = font_sec_header
            c_m.fill = fill_accent
            c_m.alignment = Alignment(horizontal="center", vertical="center")

            # Sub-headers: Giorno | Ore | Tipo
            sub_headers = ["Giorno", "Ore", "Tipo"]
            for s_idx, sh in enumerate(sub_headers):
                c_sh = ws_cal.cell(header_r + 1, col_start + s_idx, value=sh)
                c_sh.font = font_tbl_header
                c_sh.fill = fill_sub
                c_sh.alignment = Alignment(horizontal="center", vertical="center")
                c_sh.border = border_header

            # Entries for this month
            m_items = [e for e in entries if e["month_year"] == mname]
            for r_idx in range(15): # 15 rows reserved per month
                curr_row = header_r + 2 + r_idx
                if r_idx < len(m_items):
                    item = m_items[r_idx]
                    dt_val = item["date"]
                    hrs_val = item["hours"]
                    tp_val = "Maturata" if item["type"] == "maturata" else "Utilizzata"

                    c_d = ws_cal.cell(curr_row, col_start, value=dt_val)
                    c_h = ws_cal.cell(curr_row, col_start + 1, value=hrs_val)
                    c_t = ws_cal.cell(curr_row, col_start + 2, value=tp_val)

                    c_d.alignment = Alignment(horizontal="center")
                    c_h.alignment = Alignment(horizontal="right")
                    c_t.alignment = Alignment(horizontal="center")

                    if item["type"] == "utilizzata":
                        c_d.fill = fill_yellow
                        c_h.fill = fill_yellow
                        c_t.fill = fill_yellow
                    else:
                        c_d.fill = fill_green
                        c_h.fill = fill_green
                        c_t.fill = fill_green
                else:
                    c_d = ws_cal.cell(curr_row, col_start, value="")
                    c_h = ws_cal.cell(curr_row, col_start + 1, value="")
                    c_t = ws_cal.cell(curr_row, col_start + 2, value="")

                for s_idx in range(3):
                    ws_cal.cell(curr_row, col_start + s_idx).border = border_all
                    ws_cal.cell(curr_row, col_start + s_idx).font = font_regular

            # Month Total
            tot_row = header_r + 17
            ws_cal.cell(tot_row, col_start, value="Totale Ore").font = font_bold
            ws_cal.cell(tot_row, col_start).alignment = Alignment(horizontal="right")
            
            m_tot_util = sum(x["hours"] for x in m_items if x["type"] == "utilizzata")
            m_tot_mat = sum(x["hours"] for x in m_items if x["type"] == "maturata")
            m_net = m_tot_mat - m_tot_util

            c_tot = ws_cal.cell(tot_row, col_start + 1, value=m_net)
            c_tot.font = font_bold
            c_tot.number_format = '+#,##0.0;-#,##0.0;0.0'
            c_tot.alignment = Alignment(horizontal="right")

            for s_idx in range(3):
                c = ws_cal.cell(tot_row, col_start + s_idx)
                c.fill = fill_tot
                c.border = border_total

    # Adjust column widths across sheets
    for ws in [ws_dash, ws_log, ws_cal]:
        for col in ws.columns:
            max_len = 0
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len and len(val) < 40:
                    max_len = len(val)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path


def export_to_csv(employee_id: int, output_path: str) -> str:
    """Exports employee work entries to standard semicolon-delimited CSV."""
    entries = database.get_all_chronological_entries(employee_id)
    emp = database.get_employee(employee_id)
    if not emp:
        raise ValueError(f"Dipendente {employee_id} non trovato.")

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, mode="w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(["ID", "Dipendente", "Data", "Giorno", "Mese_Anno", "Ore", "Tipo", "Causale", "Note"])
        for e in entries:
            writer.writerow([
                e["id"],
                emp["name"],
                e["date"],
                e["day_of_week"],
                e["month_year"],
                e["hours"],
                e["type"],
                e["causale"],
                e["note"]
            ])
    return output_path


def import_from_excel(employee_id: int, file_path: str) -> int:
    """Imports entries from an Excel file into the employee's history."""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File non trovato: {file_path}")

    wb = openpyxl.load_workbook(file_path, data_only=True)
    imported_count = 0

    # 1. Check if 'Vista Calendario Mensile' exists
    if "Vista Calendario Mensile" in wb.sheetnames:
        ws = wb["Vista Calendario Mensile"]
        month_blocks = [
            (3, [("FEBBRAIO 2026", 2), ("MARZO 2026", 6), ("APRILE 2026", 10), ("MAGGIO 2026", 14)]),
            (23, [("GIUGNO 2026", 2), ("LUGLIO 2026", 6), ("AGOSTO 2026", 10), ("SETTEMBRE 2026", 14)]),
            (43, [("OTTOBRE 2026", 2), ("NOVEMBRE 2026", 6), ("DICEMBRE 2026", 10), ("GENNAIO 2027", 14)]),
            (63, [("FEBBRAIO 2027", 2), ("MARZO 2027", 6), ("APRILE 2027", 10), ("MAGGIO 2027", 14)])
        ]
        for header_r, row_months in month_blocks:
            for mname, c_start in row_months:
                for r in range(header_r + 2, header_r + 17):
                    giorno = ws.cell(r, c_start).value
                    ore = ws.cell(r, c_start + 1).value
                    tipo = str(ws.cell(r, c_start + 2).value or "").strip()

                    if giorno is not None and ore is not None:
                        try:
                            # Normalize date
                            if isinstance(giorno, (datetime.date, datetime.datetime)):
                                dt_str = giorno.strftime("%Y-%m-%d")
                            else:
                                raw_d = str(giorno).strip()
                                if "/" in raw_d:
                                    parts = raw_d.split("/")
                                    dt_str = f"{parts[2]}-{int(parts[1]):02d}-{int(parts[0]):02d}"
                                else:
                                    dt_str = raw_d

                            hrs = float(ore)
                            entry_type = "maturata" if "maturat" in tipo.lower() or "straordinar" in tipo.lower() else "utilizzata"

                            database.add_entry(employee_id, {
                                "date": dt_str,
                                "hours": hrs,
                                "type": entry_type,
                                "causale": f"{'Straordinario' if entry_type == 'maturata' else 'Permesso'} importato",
                                "note": ""
                            })
                            imported_count += 1
                        except Exception as ex:
                            print(f"Skipping row {r} in {mname}: {ex}")

    # 2. Or fallback to 'Registro Movimenti'
    elif "Registro Movimenti" in wb.sheetnames:
        ws = wb["Registro Movimenti"]
        for r in range(4, ws.max_row):
            data_val = ws.cell(r, 2).value
            if not data_val:
                continue
            try:
                if isinstance(data_val, (datetime.date, datetime.datetime)):
                    dt_str = data_val.strftime("%Y-%m-%d")
                else:
                    dt_str = str(data_val).strip()

                u_val = ws.cell(r, 6).value
                m_val = ws.cell(r, 7).value
                cau = str(ws.cell(r, 5).value or "")
                note = str(ws.cell(r, 9).value or "")

                if u_val and float(u_val) > 0:
                    database.add_entry(employee_id, {
                        "date": dt_str,
                        "hours": float(u_val),
                        "type": "utilizzata",
                        "causale": cau,
                        "note": note
                    })
                    imported_count += 1
                elif m_val and float(m_val) > 0:
                    database.add_entry(employee_id, {
                        "date": dt_str,
                        "hours": float(m_val),
                        "type": "maturata",
                        "causale": cau,
                        "note": note
                    })
                    imported_count += 1
            except Exception as ex:
                print(f"Skipping row {r}: {ex}")

    return imported_count


def generate_course_attendance_template(course_id: int, output_path: str) -> str:
    """
    Generates a pre-filled Excel template (.xlsx) for tutor attendance registration
    based on the course calendar sessions.
    """
    conn = database.get_db_connection()
    cur = conn.cursor()
    cur.execute("SELECT id, name, code, employee_id FROM courses WHERE id = ?", (course_id,))
    course = cur.fetchone()
    if not course:
        conn.close()
        raise ValueError(f"Corso {course_id} non trovato.")

    cur.execute("""
        SELECT s.date, s.day_of_week, s.start_time, s.end_time, s.hours, s.modulo, s.docente, s.phase
        FROM course_calendar_sessions s
        JOIN course_calendars cal ON s.calendar_id = cal.id
        WHERE cal.course_id = ?
        ORDER BY s.date ASC, s.start_time ASC, s.id ASC
    """, (course_id,))
    sessions = [dict(r) for r in cur.fetchall()]
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Presenze Corso"

    # Styling constants
    navy_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="4B5563")
    regular_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB")
    )
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    # Title block
    ws["A1"] = f"REGISTRO PRESENZE TUTOR — {course['name']} ({course['code'] or 'N/D'})"
    ws["A1"].font = title_font
    ws["A2"] = "Foglio presenze già compilato con le date del corso. Modifica gli orari effettivi se necessario e ricarica il file sul sito."
    ws["A2"].font = subtitle_font

    # Headers at row 4
    headers = [
        ("Data", 16, center_align),
        ("Giorno", 14, center_align),
        ("Ora Inizio", 12, center_align),
        ("Ora Fine", 12, center_align),
        ("Ore Presenza", 14, right_align),
        ("Modulo / Argomento Didattico", 34, left_align),
        ("Docente", 28, left_align),
        ("Fase", 12, center_align),
        ("Note", 22, left_align)
    ]

    for col_idx, (h_title, width, align) in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h_title)
        cell.font = header_font
        cell.fill = navy_fill
        cell.alignment = center_align
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    ws.row_dimensions[4].height = 26

    # Data rows
    for r_idx, s in enumerate(sessions, start=5):
        ws.row_dimensions[r_idx].height = 20
        is_even = (r_idx % 2 == 0)
        row_fill = zebra_fill if is_even else None

        # Format date as DD/MM/YYYY
        dt_str = s["date"]
        try:
            parts = dt_str.split("-")
            disp_date = f"{parts[2]}/{parts[1]}/{parts[0]}"
        except Exception:
            disp_date = dt_str

        vals = [
            (disp_date, regular_font, center_align, "@"),
            (s.get("day_of_week", ""), regular_font, center_align, "@"),
            (s.get("start_time", "09:00"), bold_font, center_align, "@"),
            (s.get("end_time", "13:00"), bold_font, center_align, "@"),
            (float(s.get("hours", 0.0)), bold_font, right_align, "0.0"),
            (s.get("modulo", "Lezione d'aula"), regular_font, left_align, "@"),
            (s.get("docente", ""), regular_font, left_align, "@"),
            ("Aula" if s.get("phase") == "aula" else "Stage", regular_font, center_align, "@"),
            ("Presenza svolta", regular_font, left_align, "@")
        ]

        for c_idx, (val, font, align, num_fmt) in enumerate(vals, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font
            cell.alignment = align
            cell.border = thin_border
            cell.number_format = num_fmt
            if row_fill:
                cell.fill = row_fill

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path

